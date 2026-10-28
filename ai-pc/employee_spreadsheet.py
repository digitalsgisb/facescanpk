"""Excel templates and parsing for previewed employee roster imports."""
import base64
import io
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [('Employee ID','id'),('Name','name'),('Plant / location','base'),
           ('Department','department'),('Line','line'),('Team','team'),
           ('Default shift ID','schedule_id'),('Nickname','nickname'),
           ('Company','company'),('Position','position'),('Office department','office_department'),
           ('Production section','production_section'),('Employment department','employment_department'),
           ('Job type','job_type')]


def template(db, plant):
    book=Workbook()
    sheet=book.active
    sheet.title='Employee import'
    sheet.append([label for label,_ in COLUMNS])
    employees=db.execute("SELECT * FROM employees WHERE active=1 AND (?='' OR base=?) ORDER BY base,name,id",(plant,plant))
    for e in employees:
        sheet.append([e[key] if e[key] not in (None,'') else 'Unassigned' if key not in ('id','name') else '' for _,key in COLUMNS])
    for row in sheet:
        for cell in row:
            cell.data_type='s'
            cell.number_format='@'
    for cell in sheet[1]:cell.font=Font(bold=True)
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    for col in 'ABCDEFGHIJKLMN':sheet.column_dimensions[col].width=25
    teams=DataValidation(type='list',formula1='"A,B,C,D,N,Unassigned"')
    sheet.add_data_validation(teams);teams.add('F2:F5001')
    info=book.create_sheet('Read me')
    for line in [
        'Edit Employee import, then upload this .xlsx file under Employees > Import employees.',
        'Employee ID identifies a person. Keep IDs as text to preserve leading zeros.',
        'Blank cells keep existing values. Type Unassigned to clear an optional field.',
        'New employees require Employee ID, Name and Plant / location. Employees absent from this file are kept.',
        'Locations use uppercase text. Team must be A, B, C, D, N or Unassigned.',
        'Default shift ID is a fallback: team rotation, group shifts and date overrides take precedence.',
        'Team/location changes automatically pick up that group’s saved shift rules.',
        'Preview every change before applying. Photos and attendance are preserved. A backup is saved before changes.',
        'This template changes employee profiles/default shifts; use Shifts for team rules or date overrides.',
        'Available default shift IDs:']:
        info.append([line])
    for schedule in db.execute('SELECT * FROM schedules ORDER BY id'):
        info.append([schedule['id'],schedule['name'],schedule['start'],schedule['end']])
    info.column_dimensions['A'].width=110
    output=io.BytesIO();book.save(output)
    return output.getvalue()


def parse(encoded, Problem):
    try:
        if not isinstance(encoded,str) or len(encoded)>2_700_000:
            raise ValueError('Choose an Excel file smaller than 2 MB.')
        raw=base64.b64decode(encoded,validate=True)
        with ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist())>200 or sum(i.file_size for i in archive.infolist())>20_000_000:
                raise ValueError('Excel file is too large when expanded.')
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except Exception:
        raise Problem('Could not read Excel file. Use the downloaded .xlsx template.')
    try:
        sheet=book['Employee import'] if 'Employee import' in book.sheetnames else book.worksheets[0]
        rows=sheet.iter_rows(max_row=5002,max_col=20)
        header=next(rows,None)
        names={label.casefold():key for label,key in COLUMNS}
        names.update({key:key for _,key in COLUMNS})
        names.update({'location':'base','plant':'base','shift id':'schedule_id','id':'id'})
        mapping={}
        for index,cell in enumerate(header or []):
            if cell.value is None:continue
            key=names.get(str(cell.value).strip().casefold())
            if not key:raise Problem('Unknown column: '+str(cell.value)+'. Use the employee import template.')
            if key in mapping.values():raise Problem('Duplicate column: '+str(cell.value))
            mapping[index]=key
        if 'id' not in mapping.values():raise Problem('Excel file needs an Employee ID column.')
        employees=[]
        for number,row in enumerate(rows,2):
            if not any(c.value is not None for c in row):continue
            if number>5001:raise Problem('Import supports up to 5000 employee rows.')
            source={}
            for index,key in mapping.items():
                cell=row[index]
                if cell.data_type=='f':raise Problem(f'Row {number}: replace formulas with plain text.')
                value=cell.value
                if value is None or str(value).strip()=='':continue
                if key=='id' and not isinstance(value,str):
                    raise Problem(f'Row {number}: format Employee ID as text to preserve leading zeros.')
                if not isinstance(value,(str,int,float)):raise Problem(f'Row {number}: use text values.')
                value=str(value).strip()
                source[key]='' if key not in ('id','name') and value.casefold()=='unassigned' else value
            if not source.get('id'):raise Problem(f'Row {number}: Employee ID is required.')
            employees.append(source)
        return {'format':'sgi-employee-roster-v1','employees':employees}
    except Problem:
        raise
    except Exception:
        raise Problem('Could not read spreadsheet rows. Use the downloaded .xlsx template.')
    finally:
        book.close()
