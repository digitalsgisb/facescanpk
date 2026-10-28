import base64
import io
import unittest
import uuid
from openpyxl import Workbook,load_workbook
import test_testing_operations as legacy


class SpreadsheetTests(unittest.TestCase):
    setUp=legacy.PlantOperationsTests.setUp
    tearDown=legacy.PlantOperationsTests.tearDown
    request=legacy.PlantOperationsTests.request
    login=legacy.PlantOperationsTests.login
    staff_login=legacy.PlantOperationsTests.staff_login

    def preview(self,headers,rows):
        book=Workbook();sheet=book.active;sheet.append(headers)
        for row in rows:sheet.append(row)
        raw=io.BytesIO();book.save(raw)
        return self.request('/api/employee-import/spreadsheet','POST',{'file_base64':base64.b64encode(raw.getvalue()).decode()})

    def apply(self,report):
        return self.request('/api/employee-import/apply','POST',{'payload':report['payload'],'preview_token':report['preview_token'],'request_id':str(uuid.uuid4())})

    def test_bulk_changes_are_previewed_then_applied_without_losing_photos(self):
        status,report=self.preview(['Employee ID','Plant / location','Team','Default shift ID','Line'],[['PK','SENDAYAN','C','night-ot','ROLL MAKING']])
        self.assertEqual(status,200)
        self.assertEqual(self.request('/api/employees')[1]['employees'][0]['team'],'N')
        self.assertEqual(self.apply(report)[0],200)
        people={e['id']:e for e in self.request('/api/employees')[1]['employees']}
        self.assertEqual([people['PK'][k] for k in ('base','team','schedule_id','line')],['SENDAYAN','C','night-ot','ROLL MAKING'])
        self.assertEqual(people['PK']['photo_url'],'https://example.invalid/photo.jpg')
        self.assertEqual(people['SD']['team'],'N')

    def test_template_round_trip_has_no_changes(self):
        status,raw=self.request('/api/employee-template')
        self.assertEqual(status,200)
        self.assertIn('Read me',load_workbook(io.BytesIO(raw)).sheetnames)
        status,report=self.request('/api/employee-import/spreadsheet','POST',{'file_base64':base64.b64encode(raw).decode()})
        self.assertEqual(status,200)
        self.assertEqual(report['summary']['unchanged'],2)

    def test_blank_preserves_unassigned_clears_and_new_employee_is_added(self):
        status,report=self.preview(['Employee ID','Name','Plant / location','Team','Default shift ID'],
                                   [['PK',None,None,'Unassigned','Unassigned'],['NEW','New Person','SENDAYAN','D','day']])
        self.assertEqual(status,200)
        self.assertEqual(self.apply(report)[0],200)
        people={e['id']:e for e in self.request('/api/employees')[1]['employees']}
        self.assertEqual(people['PK']['name'],'PK')
        self.assertEqual(people['PK']['team'],'')
        self.assertIsNone(people['PK']['schedule_id'])
        self.assertEqual(people['NEW']['team'],'D')

    def test_invalid_duplicate_formula_numeric_id_and_cross_plant_rows(self):
        for rows in ([['PK','Z']],[['PK','A'],['PK','B']],[[123,'A']],[['PK','=1+1']]):
            self.assertEqual(self.preview(['Employee ID','Team'],rows)[0],400)
        self.assertEqual(self.preview(['Employee ID','Default shift ID'],[['PK','missing']])[0],400)
        self.staff_login()
        self.assertEqual(self.preview(['Employee ID','Team'],[['SD','C']])[0],403)
        self.assertEqual(self.preview(['Employee ID','Plant / location'],[['PK','SENDAYAN']])[0],403)

    def test_stale_preview_cannot_overwrite_new_edits(self):
        status,report=self.preview(['Employee ID','Team'],[['PK','B']])
        self.assertEqual(status,200)
        self.request('/api/employees/PK','PUT',{'nickname':'New name'})
        self.assertEqual(self.apply(report)[0],409)
